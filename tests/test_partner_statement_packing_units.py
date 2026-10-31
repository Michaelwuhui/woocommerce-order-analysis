"""Per-unit contract regressions; all orders and shipments are synthetic."""
from copy import deepcopy
from io import BytesIO
import json

from openpyxl import load_workbook
import pytest

from partner_statement_sources import read_version, save_version
from test_partner_statement_sources import db, order, assign, parcel, snapshot, routes


def unit_rule(db, **changes):
    rule={'partner_id':1,'warehouse_id':10,'billing_basis':'shipped_unit','unit_fee_cny':'10.00',
          'effective_from':'2026-09-01','effective_to':'2099-01-01',**changes}
    db.execute('INSERT INTO rec_objects(id,kind,name,data,created_at) VALUES(?,?,?,?,?)',
               ('unit-rule','partner_packing_rule','逐支打包费',json.dumps(rule),'2026-10-31'))


@pytest.mark.parametrize('status,returned',[
    ('completed',False),('failed',False),('cancelled',False),('refunded',False),('on-hold',True)])
def test_each_physically_shipped_transit_unit_is_charged_regardless_of_order_outcome(db,status,returned):
    unit_rule(db)
    order(db,status=status,returned=returned,items=[{'id':11,'name':'Transit','quantity':3,'total':'300'}])
    item=assign(db,'o1',11,10,qty=3);parcel(db,'o1','sent',[item],logs=False)
    db.execute('UPDATE oms_shipment_items SET quantity=3')
    value=snapshot(db);fee=value['service_lines'][0]
    assert value['settlement']['packing_fee_cny']==30
    assert value['settlement']['packing_units']==3
    assert fee['quantity']==3 and fee['unit_fee_cny']==10 and fee['order_status']==status
    assert fee['items'][0]['line_item_id']=='11'
    assert not fee['pending']


def test_joint_parcel_only_charges_its_transit_goods_and_counts_reships_including_omissions(db):
    unit_rule(db);order(db)
    regular=assign(db,'o1',11,1);transit=assign(db,'o1',12,10)
    parcel(db,'o1','joint',[regular,transit],primary=1)
    parcel(db,'o1','reship',[transit],reship=True,date='2026-09-12')
    value=snapshot(db,decisions={'reship':{'reason':'warehouse_omission','evidence':'合约已改为实际发出逐支收费'}})
    assert value['settlement']['packing_fee_cny']==20
    assert value['settlement']['packing_units']==2
    assert all(row['quantity']==1 and len(row['items'])==1 for row in value['service_lines'])
    assert all(row['items'][0]['line_item_id']=='12' for row in value['service_lines'])
    assert value['settlement']['ready']


def test_split_parcel_is_billed_by_its_own_units_in_actual_dispatch_month(db):
    unit_rule(db);order(db,date='2026-08-10')
    a=assign(db,'o1',11,10);b=assign(db,'o1',12,10)
    parcel(db,'o1','aug',[a],date='2026-08-20',logs=False)
    parcel(db,'o1','sep',[b],date='2026-09-02',logs=False)
    value=snapshot(db)
    assert value['source_lines']==[]
    assert value['settlement']['packing_units']==1
    assert value['settlement']['packing_fee_cny']==10
    assert [r['shipment_id'] for r in value['service_lines']]==['sep']


@pytest.mark.parametrize('warehouse,confirmed_mode,expected',[(1,'transit_service',10),(10,'partner_stock',0)])
def test_packing_uses_confirmed_product_ownership_instead_of_parcel_primary_warehouse(db,warehouse,confirmed_mode,expected):
    unit_rule(db);order(db);item=assign(db,'o1',11,warehouse)
    parcel(db,'o1','sent',[item],primary=warehouse)
    proof={'order_id':'o1','line_item_id':'11','product_id':101,'variation_id':0,'quantity':1,
           'mode':confirmed_mode,'evidence':'逐商品确认货权'}
    db.execute('INSERT INTO rec_objects(id,kind,name,data,created_at) VALUES(?,?,?,?,?)',
               ('source-proof','legacy_product_source','货权',json.dumps(proof),'2026-10-31'))
    assert snapshot(db)['settlement']['packing_fee_cny']==expected
    db.execute("UPDATE orders SET line_items=replace(line_items,'101','999')")
    value=snapshot(db)
    assert value['settlement']['packing_pending']==1
    assert value['service_lines'][0]['amount_cny'] is None


def test_duplicate_physical_tracking_merges_identical_items_without_double_charge(db):
    unit_rule(db);order(db);a=assign(db,'o1',11,10)
    parcel(db,'o1','one',[a]);parcel(db,'o1','duplicate',[a])
    db.execute("UPDATE oms_shipments SET tracking_number='same-tracking'")
    value=snapshot(db)
    assert value['settlement']['packing_fee_cny']==10
    assert len(value['service_lines'])==1
    assert value['service_lines'][0]['shipment_ids']==['duplicate','one']
    db.execute("UPDATE oms_shipment_items SET quantity=2 WHERE shipment_id='duplicate'")
    value=snapshot(db)
    assert value['settlement']['packing_pending']==1
    assert value['service_lines'][0]['amount_cny'] is None
    assert '同一运单商品数量冲突' in value['service_lines'][0]['pending']


@pytest.mark.parametrize('missing', ['tracking','items','invalid_quantity'])
def test_missing_physical_quantity_or_tracking_is_unknown_not_a_zero_fee(db,missing):
    unit_rule(db);order(db);a=assign(db,'o1',11,10);parcel(db,'o1','sent',[a])
    if missing=='tracking':db.execute("UPDATE oms_shipments SET tracking_number=' '")
    elif missing=='items':db.execute('DELETE FROM oms_shipment_items')
    else:db.execute('UPDATE oms_shipment_items SET quantity=0')
    value=snapshot(db)
    assert value['settlement']['packing_pending']==1
    assert value['service_lines'][0]['amount_cny'] is None


def test_label_only_is_not_billed_but_cancellation_after_real_carrier_delivery_is_billed(db):
    unit_rule(db);order(db);a=assign(db,'o1',11,10);parcel(db,'o1','sent',[a])
    db.execute("UPDATE oms_shipments SET status='label_pending'")
    assert snapshot(db)['service_lines']==[]
    db.execute("UPDATE oms_shipments SET status='cancelled'")
    assert snapshot(db)['service_lines']==[]
    db.execute('CREATE TABLE oms_tracking_events(shipment_id TEXT,normalized_status TEXT)')
    db.execute("INSERT INTO oms_tracking_events VALUES('sent','delivered')")
    assert snapshot(db)['settlement']['packing_fee_cny']==10


def test_new_contract_has_its_own_scope_and_does_not_rewrite_frozen_old_bill(db):
    order(db);a=assign(db,'o1',11,10);parcel(db,'o1','sent',[a])
    save_version(db,59,snapshot(db),1)
    before=deepcopy(read_version(db,59,details=True))
    unit_rule(db,partner_id=99)
    assert snapshot(db)['settlement']['packing_fee_cny']==20
    rule=json.loads(db.execute("SELECT data FROM rec_objects WHERE id='unit-rule'").fetchone()[0])
    rule['partner_id']=1
    db.execute("UPDATE rec_objects SET data=? WHERE id='unit-rule'",(json.dumps(rule),))
    after=snapshot(db)
    assert after['settlement']['packing_fee_cny']==10
    assert read_version(db,59,details=True)==before
    save_version(db,59,after,1)
    assert read_version(db,59,revision=1,details=True)==before


def test_empty_month_still_displays_per_unit_contract(db):
    unit_rule(db);order(db,warehouse=10)
    value=snapshot(db)
    assert value['settlement']['packing_basis']=='shipped_unit'
    assert value['settlement']['packing_units']==0
    assert value['settlement']['packing_fee_cny']==0
    assert value['settlement']['packing_unit_fee_cny']==10


def test_per_unit_export_and_search_trace_the_same_frozen_shipment_items(db):
    unit_rule(db);order(db);regular=assign(db,'o1',11,1);transit=assign(db,'o1',12,10)
    parcel(db,'o1','joint',[regular,transit],primary=1,logs=False)
    value=snapshot(db)
    db.execute("INSERT INTO reconciliation_statements(id,partner_id,status,period_year,period_month) VALUES(59,1,'generated',2026,9)")
    save_version(db,59,value,1);db.commit()
    app,_=routes(db);client=app.test_client()
    result=client.get('/api/reconciliation/statements/59/source-detail?kind=fees&search=Item%20B').json
    assert result['count']==1 and result['rows'][0]['quantity']==1
    book=load_workbook(BytesIO(client.get('/api/reconciliation/statements/59/export').data),data_only=True)
    fee=list(book['中转打包费明细'].iter_rows(values_only=True))
    assert fee[1][8]==10 and fee[1][13:16]==(1,10,'实发支数')
    items=list(book['打包商品明细'].iter_rows(values_only=True))
    assert len(items)==2 and items[1][8]=='Item B' and items[1][10:13]==(1,10,10)
    basis=dict(book['版本与口径'].iter_rows(min_row=2,values_only=True))
    assert '失败' in basis['打包费口径'] and '逐支收费' in basis['补发口径']
