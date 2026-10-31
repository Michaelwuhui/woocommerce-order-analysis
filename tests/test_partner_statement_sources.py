import ast
from copy import deepcopy
from decimal import Decimal
from io import BytesIO
import json
from pathlib import Path
import sqlite3
from types import SimpleNamespace

from flask import Flask, jsonify, request, send_file
from openpyxl import load_workbook
import pytest

from partner_statement_sources import allocate, build_snapshot, init_sqlite_schema, read_version, save_version


class Connection(sqlite3.Connection):
    statements = 0
    def execute(self,*args,**kwargs):
        self.statements += 1
        return super().execute(*args,**kwargs)
    def close(self):
        pass


@pytest.fixture
def db():
    c = sqlite3.connect(':memory:',factory=Connection)
    c.row_factory = sqlite3.Row
    c.executescript('''
        CREATE TABLE partners(id INTEGER PRIMARY KEY,name TEXT,code TEXT,currency TEXT,cost_ratio REAL,partner_profit_ratio REAL,our_profit_ratio REAL);
        INSERT INTO partners VALUES(1,'合作仓','poland_jin','PLN',.5,.25,.25);
        CREATE TABLE sites(id INTEGER PRIMARY KEY,url TEXT,country TEXT,manager TEXT,cod_on_hold_is_shipped INTEGER);
        INSERT INTO sites VALUES(1,'https://shop.test','PL','Alice',1);
        INSERT INTO sites VALUES(2,'https://outside.test','PL','Bob',1);
        CREATE TABLE effective_partner_sites(partner_id INTEGER,site_id INTEGER);
        INSERT INTO effective_partner_sites VALUES(1,1);
        CREATE TABLE warehouses(id INTEGER PRIMARY KEY,code TEXT,name TEXT,country TEXT);
        INSERT INTO warehouses VALUES(1,'PL','自备货仓','PL');
        INSERT INTO warehouses VALUES(10,'PL-JYJG-TRANSIT','中转仓','PL');
        CREATE TABLE orders(id TEXT PRIMARY KEY,number TEXT,source TEXT,status TEXT,payment_method TEXT,total REAL,shipping_total REAL,
            date_created TEXT,currency TEXT,warehouse_id INTEGER,line_items TEXT,is_undelivered INTEGER,is_problem_return INTEGER,shipping_loss_amount REAL);
        CREATE TABLE oms_order_items(id INTEGER PRIMARY KEY,woo_line_item_id TEXT,line_index INTEGER);
        CREATE TABLE oms_fulfillments(id TEXT PRIMARY KEY,order_id TEXT,warehouse_id INTEGER,status TEXT);
        CREATE TABLE oms_fulfillment_items(id INTEGER PRIMARY KEY,fulfillment_id TEXT,order_item_id INTEGER,allocated_qty INTEGER,fulfilled_qty INTEGER,cancelled_qty INTEGER);
        CREATE TABLE oms_shipments(id TEXT PRIMARY KEY,fulfillment_id TEXT,tracking_number TEXT,shipped_at TEXT,status TEXT);
        CREATE TABLE oms_shipment_items(shipment_id TEXT,fulfillment_item_id INTEGER,quantity INTEGER);
        CREATE TABLE shipping_logs(id INTEGER PRIMARY KEY,order_id TEXT,tracking_number TEXT,is_reship INTEGER,reship_reason TEXT,shipped_at TEXT,status TEXT);
        CREATE TABLE rec_objects(id TEXT,kind TEXT,name TEXT,data TEXT,created_at TEXT,actor TEXT);
        CREATE TABLE reconciliation_statements(id INTEGER PRIMARY KEY,partner_id INTEGER,period_year INTEGER,period_month INTEGER,
            total_orders INTEGER,total_gross_pln REAL,total_net_pln REAL,cost_amount_pln REAL,partner_profit_pln REAL,our_receivable_pln REAL,
            exchange_rate_cny REAL,our_receivable_cny REAL,actual_cost_pln_snapshot REAL,calc_mode TEXT,status TEXT,
            is_manual INTEGER DEFAULT 0,notes TEXT,created_at TEXT DEFAULT CURRENT_TIMESTAMP,updated_at TEXT DEFAULT CURRENT_TIMESTAMP,
            confirmed_at TEXT,confirmed_by INTEGER,confirmed_name TEXT,confirmed_ip TEXT,locked_at TEXT);
        CREATE TABLE reconciliation_statement_orders(statement_id INTEGER,order_id TEXT,order_number TEXT,status_at_gen TEXT,total_at_gen REAL,
            shipping_at_gen REAL,shipping_loss_at_gen REAL,is_undelivered_at_gen INTEGER,currency_at_gen TEXT,date_created TEXT);
        CREATE TABLE partner_receipts(id INTEGER PRIMARY KEY,statement_id INTEGER,amount_pln REAL,receipt_date TEXT);
    ''')
    rule={'warehouse_id':10,'management_cny':'20.00','provider_id':'provider',
          'effective_from':'2026-08-01','effective_to':'2099-01-01'}
    c.execute('INSERT INTO rec_objects(id,kind,name,data,created_at) VALUES (?,?,?,?,?)',('rule','history_rule','打包费',json.dumps(rule),'2026-08-01'))
    init_sqlite_schema(c)
    yield c
    sqlite3.Connection.close(c)


def order(db, ident='o1', *, items=None, warehouse=None, status='completed', returned=False, loss=0, date='2026-09-10'):
    if items is None:
        items=[{'id':11,'name':'Item A','quantity':1,'total':'100','product_id':101},
               {'id':12,'name':'Item B','quantity':1,'total':'100','product_id':102}]
    net=sum(Decimal(str(item['total'])) for item in items)
    db.execute('INSERT INTO orders VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
        (ident,ident,'https://shop.test',status,'cod',float(net+25),25,date,'PLN',warehouse,json.dumps(items),returned,0,loss))
    return items


def assign(db, ident, line_id, warehouse, qty=1, *, status='delivered', fulfilled=None):
    fiid = db.execute('SELECT COALESCE(MAX(id),0)+1 FROM oms_fulfillment_items').fetchone()[0]
    oid=fiid
    db.execute('INSERT INTO oms_order_items VALUES (?,?,?)',(oid,str(line_id),0))
    fid=ident+':'+str(warehouse)
    db.execute('INSERT OR IGNORE INTO oms_fulfillments VALUES (?,?,?,?)',(fid,ident,warehouse,status))
    db.execute('INSERT INTO oms_fulfillment_items VALUES (?,?,?,?,?,?)',(fiid,fid,oid,qty,qty if fulfilled is None else fulfilled,0))
    return fiid


def parcel(db, ident, sid, parts, *, primary=10, date='2026-09-11', reship=False, logs=True):
    db.execute('INSERT INTO oms_shipments VALUES (?,?,?,?,?)',(sid,ident+':'+str(primary),'TRACK-'+sid,date,'shipped'))
    for part in parts:
        db.execute('INSERT INTO oms_shipment_items VALUES (?,?,?)',(sid,part,1))
    if logs:
        db.execute('INSERT INTO shipping_logs(order_id,tracking_number,is_reship,reship_reason,shipped_at,status) VALUES (?,?,?,?,?,?)',
                   (ident,'TRACK-'+sid,reship,'',date,'shipped'))


def snapshot(db, *, mode='actual', decisions=None, lookup=None, month=9):
    partner=dict(db.execute('SELECT * FROM partners WHERE id=1').fetchone())
    business={'total_orders':1,'total_gross_pln':225,'total_net_pln':200,'actual_cost_pln':80,
              'by_status':{},'by_product':[],'by_site':[]}
    return build_snapshot(db,partner,2026,month,mode,business,
        lookup or (lambda *args:{'unit_cost':40,'cost_id':7,'effective_date':'2026-08-01'}),decisions=decisions)


@pytest.mark.parametrize('amount,weights',[(.05,[1]*10),(-.05,[1]*10),(.01,[0,0,0]),(100,[1,3,9]),(-2,[1,0,1])])
def test_cent_allocation_conserves_money_without_negative_positive_shares(amount,weights):
    values=allocate(amount,weights)
    assert sum(values)==Decimal(str(amount))
    assert all(value>=0 for value in values) if amount>=0 else all(value<=0 for value in values)


def test_mixed_order_keeps_only_partner_goods_margin_and_charges_joint_parcel_once(db):
    order(db); regular=assign(db,'o1',11,1); transit=assign(db,'o1',12,10)
    parcel(db,'o1','joint',[regular,transit],primary=1)
    result=snapshot(db)
    assert result['settlement']['cost_recovery']==40
    assert result['settlement']['partner_profit']==30
    assert result['settlement']['our_receivable']==130
    assert result['settlement']['net']==200
    assert result['settlement']['packing_fee_cny']==20
    assert result['settlement']['ready'] is True
    middle=next(row for row in result['source_lines'] if row['mode']=='transit_service')
    assert (middle['cost_recovery'],middle['partner_profit'],middle['our_receivable'])==(0,0,100)
    assert result['service_lines'][0]['shipment_id']=='joint'


def test_same_product_split_between_suppliers_is_split_by_quantity(db):
    order(db,items=[{'id':11,'name':'Same item','quantity':2,'total':'200'}])
    assign(db,'o1',11,1);assign(db,'o1',11,10)
    result=snapshot(db)
    assert len(result['source_products'])==2
    assert [row['revenue'] for row in result['source_lines']]==[100,100]
    assert result['settlement']['our_receivable']==130


def test_period_profit_rounding_does_not_give_all_half_cents_to_one_party(db):
    for index in range(10):
        order(db,str(index),warehouse=1,items=[{'id':11,'name':'Same item','quantity':1,'total':'100.01'}])
    result=snapshot(db)
    assert result['settlement']['partner_profit']==300.05
    assert result['settlement']['our_receivable']==300.05
    assert sum(Decimal(str(row['partner_profit'])) for row in result['source_lines'])==Decimal('300.05')
    for row in result['source_lines']:
        assert sum(Decimal(str(row[key])) for key in ('cost_recovery','partner_profit','our_receivable'))==Decimal(str(row['revenue']))


def test_transit_never_looks_up_or_refunds_supplier_cost(db):
    order(db,warehouse=10)
    def forbidden(*args):raise AssertionError('Transit stock is not partner capital')
    result=snapshot(db,lookup=forbidden)
    assert result['settlement']['cost_recovery']==0
    assert result['settlement']['partner_profit']==0
    assert result['settlement']['our_receivable']==200


def test_explicit_ownership_confirmation_is_per_order_item_and_keeps_physical_warehouse(db):
    order(db,warehouse=1)
    proof={'order_id':'o1','line_item_id':'11','product_id':101,'variation_id':0,'quantity':1,
           'mode':'transit_service','evidence':'业务负责人确认本件为我方备货'}
    db.execute('INSERT INTO rec_objects(id,kind,name,data,created_at) VALUES(?,?,?,?,?)',('confirmation','legacy_product_source','供货确认',json.dumps(proof),'2026-10-31'))
    result=snapshot(db)
    first,second=result['source_lines']
    assert first['mode']=='transit_service' and first['warehouse_id']==1
    assert first['source_override_id']=='confirmation'
    assert first['cost_recovery']==0 and first['partner_profit']==0
    assert second['mode']=='partner_stock'
    assert result['settlement']['packing_fee_cny']==0  # no actual shipping evidence
    db.execute("UPDATE orders SET line_items=replace(line_items,'101','999')")
    assert snapshot(db)['source_lines'][0]['mode']=='pending'


def test_return_keeps_service_fee_but_no_goods_profit_or_capital(db):
    order(db,returned=True,loss=50)
    a=assign(db,'o1',11,10);b=assign(db,'o1',12,10);parcel(db,'o1','ret',[a,b])
    result=snapshot(db)
    assert result['settlement']['packing_fee_cny']==20
    assert result['settlement']['transit_shipping_loss_separate']==50
    assert result['settlement']['our_receivable']==0
    assert result['settlement']['partner_profit']==0
    assert result['settlement']['cost_recovery']==0


def test_regular_return_shares_loss_without_reducing_recovered_goods_capital(db):
    order(db,warehouse=1,returned=True,loss=50)
    result=snapshot(db)
    assert result['settlement']['partner_profit']==-25
    assert result['settlement']['our_receivable']==-25
    assert result['settlement']['cost_recovery']==0


def test_missing_cost_is_pending_not_zero_profit_calculation(db):
    order(db);assign(db,'o1',11,1);assign(db,'o1',12,10)
    result=snapshot(db,lookup=lambda *args:None)
    pending=next(row for row in result['source_lines'] if row['mode']=='partner_stock')
    assert pending['cost_recovery'] is None and pending['partner_profit'] is None
    assert result['settlement']['ready'] is False
    assert result['settlement']['pending_orders']==1
    assert result['settlement']['our_receivable']==100


def test_missing_oms_assignment_is_pending_and_superseded_sources_never_count(db):
    order(db);assign(db,'o1',11,10)
    old=assign(db,'o1',12,1,status='superseded')
    result=snapshot(db)
    assert next(row for row in result['source_lines'] if row['line_item_id']=='12')['mode']=='pending'
    assert not any(row.get('fulfillment_item_id')==old for row in result['source_lines'])


def test_monthly_fee_uses_first_shipped_original_order_not_order_month_or_split_count(db):
    order(db,date='2026-08-20');a=assign(db,'o1',11,10);b=assign(db,'o1',12,10)
    parcel(db,'o1','first',[a],date='2026-08-25');parcel(db,'o1','split',[b],date='2026-09-02')
    result=snapshot(db)
    assert result['source_lines']==[]
    assert result['settlement']['packing_fee_cny']==0
    assert result['service_lines'][0]['reason']=='同一原订单拆包不重复收费'


def test_original_shipped_this_month_from_previous_order_month_has_service_fee(db):
    order(db,date='2026-08-20');a=assign(db,'o1',11,10);parcel(db,'o1','first',[a])
    assert snapshot(db)['settlement']['packing_fee_cny']==20


def test_unknown_prior_parcel_cannot_make_next_month_split_charge_as_new_original(db):
    order(db,date='2026-08-20');a=assign(db,'o1',11,10);b=assign(db,'o1',12,10)
    parcel(db,'o1','first',[a],date='2026-08-25',logs=False)
    parcel(db,'o1','split',[b],date='2026-09-02')
    assert snapshot(db)['settlement']['packing_pending']==1
    classification={'shipment_id':'first','order_id':'o1','reason':'original','evidence':'核对原始出库凭证'}
    db.execute('INSERT INTO rec_objects(id,kind,name,data,created_at) VALUES(?,?,?,?,?)',('classification','legacy_shipment_classification','出库分类',json.dumps(classification),'2026-10-31'))
    result=snapshot(db)
    assert result['settlement']['packing_pending']==0
    assert result['settlement']['packing_fee_cny']==0


def test_missing_fee_amount_is_pending_while_explicit_zero_is_valid(db):
    order(db);a=assign(db,'o1',11,10);parcel(db,'o1','first',[a])
    rule=json.loads(db.execute("SELECT data FROM rec_objects WHERE id='rule'").fetchone()[0])
    rule['management_cny']=None
    db.execute("UPDATE rec_objects SET data=? WHERE id='rule'",(json.dumps(rule),))
    assert snapshot(db)['service_lines'][0]['amount_cny'] is None
    rule['management_cny']='0'
    db.execute("UPDATE rec_objects SET data=? WHERE id='rule'",(json.dumps(rule),))
    assert snapshot(db)['service_lines'][0]['amount_cny']==0


def test_late_entered_shipping_log_still_identifies_the_original_parcel(db):
    order(db);a=assign(db,'o1',11,10);parcel(db,'o1','first',[a])
    db.execute("UPDATE shipping_logs SET shipped_at='2026-10-01'")
    result=snapshot(db)
    assert result['settlement']['packing_fee_cny']==20
    assert result['settlement']['packing_pending']==0


def test_source_change_preserves_legacy_revenue_status_policy(db):
    order(db,warehouse=1)
    db.execute('UPDATE orders SET is_problem_return=1')
    assert snapshot(db)['groups'][0]['revenue']==200


@pytest.mark.parametrize('category,expected',[('ordinary',40),('warehouse_omission',20)])
def test_confirmed_reship_policy_charges_ordinary_and_waives_warehouse_omission(db,category,expected):
    order(db);a=assign(db,'o1',11,10);assign(db,'o1',12,10);parcel(db,'o1','first',[a]);parcel(db,'o1','again',[a],reship=True,date='2026-09-12')
    unresolved=snapshot(db)
    assert unresolved['settlement']['packing_pending']==1
    result=snapshot(db,decisions={'again':{'reason':category,'evidence':'核对原始发货凭证'}})
    assert result['settlement']['packing_fee_cny']==expected
    assert result['settlement']['ready'] is True
    with pytest.raises(ValueError,match='依据'):
        snapshot(db,decisions={'again':{'reason':category}})


def test_duplicate_consistent_shipping_logs_do_not_double_charge(db):
    order(db);a=assign(db,'o1',11,10);parcel(db,'o1','first',[a])
    db.execute("INSERT INTO shipping_logs(order_id,tracking_number,is_reship,shipped_at,status) SELECT order_id,tracking_number,is_reship,shipped_at,status FROM shipping_logs")
    result=snapshot(db)
    assert result['settlement']['packing_fee_cny']==20
    assert result['settlement']['packing_pending']==0


def test_conflicting_logs_require_explicit_evidence_and_preserve_classification(db):
    order(db);a=assign(db,'o1',11,10);parcel(db,'o1','first',[a],logs=False)
    result=snapshot(db)
    assert result['settlement']['packing_pending']==1
    result=snapshot(db,decisions={'first':{'reason':'original','evidence':'核实原始出库单'}})
    assert result['settlement']['packing_fee_cny']==20
    assert result['service_lines'][0]['decision']['reason']=='original'


def test_frozen_versions_remain_unchanged_when_live_orders_change(db):
    order(db,warehouse=10);result=snapshot(db)
    save_version(db,59,result,1,{'statement':{'our_receivable_pln':9},'orders':[{'order_id':'old'}]})
    expected=deepcopy(read_version(db,59,details=True))
    db.execute("UPDATE orders SET total=1000,status='cancelled'")
    assert read_version(db,59,details=True)==expected
    save_version(db,59,snapshot(db),1,{'statement':{'our_receivable_pln':200}})
    assert read_version(db,59,details=True,revision=1)==expected
    assert read_version(db,59)['revision']==2


def routes(db):
    app=Flask('source-bill-tests');app.secret_key='synthetic-only'
    user=SimpleNamespace(id=1,is_authenticated=True,get_accessible_partner_ids=lambda:None)
    ns={'app':app,'login_required':lambda f:f,'reconciliation_api_required':lambda f:f,
        'request':request,'jsonify':jsonify,'send_file':send_file,'json':json,'current_user':user,
        'sqlite3':sqlite3,'datetime':__import__('datetime').datetime,
        'get_db_connection':lambda:db,'_check_partner_access':lambda pid:pid==1,
        '_is_reconciliation_admin':lambda:True,'_enrich_statement_cny':lambda result,curr:result.update(effective_rate_cny=2),
        '_audit_log':lambda *args,**kwargs:None}
    names={'api_get_statement','api_statement_source_detail','api_statement_source_versions',
           'api_get_statement_snapshot_orders','api_export_statement','api_update_statement','api_generate_statement'}
    tree=ast.parse((Path(__file__).resolve().parents[1]/'app.py').read_text(encoding='utf-8'))
    selected=[node for node in tree.body if isinstance(node,ast.FunctionDef) and node.name in names]
    exec(compile(ast.Module(body=selected,type_ignores=[]),'app_api_subset','exec'),ns)
    return app,ns


def saved(db,result=None):
    order(db,warehouse=10)
    result=result or snapshot(db)
    db.execute("INSERT INTO reconciliation_statements(id,partner_id,period_year,period_month,total_orders,total_gross_pln,total_net_pln,cost_amount_pln,partner_profit_pln,our_receivable_pln,actual_cost_pln_snapshot,calc_mode,status) VALUES(59,1,2026,9,1,225,200,0,0,200,0,'actual','generated')")
    save_version(db,59,result,1)
    db.commit()
    return result


def test_bill_open_reads_frozen_summary_without_live_aggregation(db):
    result=saved(db);app,ns=routes(db)
    ns['_calc_partner_recon_detail']=lambda *a,**k:pytest.fail('Bill open must not recalculate live orders')
    response=app.test_client().get('/api/reconciliation/statements/59')
    assert response.status_code==200
    assert response.json['source_snapshot']['settlement']==result['settlement']
    assert 'source_lines' not in response.json['source_snapshot']


def test_saved_line_pagination_search_and_access_are_enforced(db):
    result=saved(db);app,ns=routes(db);client=app.test_client()
    response=client.get('/api/reconciliation/statements/59/source-detail?kind=lines&per_page=1&page=2&search=Item')
    assert response.json['count']==2 and len(response.json['rows'])==1
    assert response.json['rows'][0]['product_name']=='Item B'
    ns['_check_partner_access']=lambda pid:False
    for suffix in ('','/source-detail','/versions','/export','/snapshot-orders'):
        assert client.get('/api/reconciliation/statements/59'+suffix).status_code==403


def test_export_uses_same_frozen_items_and_separate_fee_currency(db):
    result=saved(db);db.execute("UPDATE orders SET total=1000,status='cancelled'")
    app,ns=routes(db)
    ns['_calc_partner_recon_detail']=lambda *a,**k:pytest.fail('Export must not recalculate')
    response=app.test_client().get('/api/reconciliation/statements/59/export')
    assert response.status_code==200
    book=load_workbook(BytesIO(response.data),data_only=True)
    assert {'按产品来源','商品来源明细','中转打包费明细','版本与口径'}<=set(book.sheetnames)
    assert book['商品来源明细'].cell(2,12).value==100
    assert book['商品来源明细'].cell(2,13).value==0
    assert book['商品来源明细'].cell(2,14).value==0


def test_incomplete_bill_cannot_be_confirmed_or_locked(db):
    order(db,warehouse=1);result=snapshot(db,lookup=lambda *a:None)
    db.execute("INSERT INTO reconciliation_statements(id,partner_id,status) VALUES(59,1,'generated')")
    save_version(db,59,result,1);db.commit()
    app,ns=routes(db);client=app.test_client()
    for action in ('confirm','lock','settle'):
        response=client.put('/api/reconciliation/statements/59',json={'action':action})
        assert response.status_code==400 and '待核实' in response.json['error']
    assert db.execute('SELECT status FROM reconciliation_statements WHERE id=59').fetchone()[0]=='generated'


def test_current_snapshot_drift_is_one_join_not_per_order_queries(db):
    saved(db)
    for index in range(50):
        db.execute('INSERT INTO reconciliation_statement_orders VALUES (?,?,?,?,?,?,?,?,?,?)',(59,'o1','o1','completed',225,25,0,0,'PLN','2026-09-10'))
    app,ns=routes(db);db.statements=0
    response=app.test_client().get('/api/reconciliation/statements/59/snapshot-orders?per_page=10&page=2')
    assert response.json['count']==50 and len(response.json['orders'])==10
    assert db.statements==2


@pytest.mark.parametrize('status', ['locked','settled','confirmed','disputed'])
def test_regeneration_cannot_overwrite_reviewed_or_locked_bills(db,status):
    saved(db);db.execute('UPDATE reconciliation_statements SET status=?',(status,));db.commit()
    app,ns=routes(db)
    response=app.test_client().post('/api/reconciliation/statements/generate',json={'partner_id':1,'year':2026,'month':9,'mode':'actual'})
    assert response.status_code==400
    assert db.execute('SELECT COUNT(*) FROM reconciliation_statement_versions').fetchone()[0]==1


def test_regeneration_cannot_overwrite_receipted_bills(db):
    saved(db);db.execute("INSERT INTO partner_receipts(statement_id,amount_pln) VALUES(59,1)");db.commit()
    app,ns=routes(db)
    response=app.test_client().post('/api/reconciliation/statements/generate',json={'partner_id':1,'year':2026,'month':9,'mode':'actual'})
    assert response.status_code==400
    assert db.execute('SELECT COUNT(*) FROM reconciliation_statement_versions').fetchone()[0]==1


def test_regeneration_is_atomic_and_preserves_previous_amounts_and_order_list(db):
    original=saved(db)
    db.execute("INSERT INTO reconciliation_statement_orders VALUES(59,'o1','o1','completed',225,25,0,0,'PLN','2026-09-10')")
    db.commit()
    app,ns=routes(db)
    ns['_calc_partner_recon_detail']=lambda *a,**k:{'total_orders':1,'total_gross_pln':225,'total_net_pln':200,'actual_cost_pln':80}
    ns['_compute_statement_split']=lambda *a,**k:(80,60,60)
    ns['_lookup_partner_rate']=lambda *a:(2,'2026-09')
    ns['_build_product_source_statement']=lambda *a,**k:original
    ns['_snapshot_statement_orders']=lambda *a:1
    response=app.test_client().post('/api/reconciliation/statements/generate',json={'partner_id':1,'year':2026,'month':9,'mode':'actual'})
    assert response.status_code==200,response.json
    assert response.json['revision']==2
    stored=json.loads(db.execute('SELECT previous_json FROM reconciliation_statement_versions WHERE revision=2').fetchone()[0])
    assert stored['statement']['our_receivable_pln']==200
    assert stored['orders'][0]['order_id']=='o1'
    def fail(*args):raise ValueError('simulated snapshot failure')
    ns['_snapshot_statement_orders']=fail
    response=app.test_client().post('/api/reconciliation/statements/generate',json={'partner_id':1,'year':2026,'month':9,'mode':'actual'})
    assert response.status_code==400
    assert db.execute('SELECT MAX(revision) FROM reconciliation_statement_versions').fetchone()[0]==2


def test_explicit_reship_classification_is_saved_with_the_generated_version(db):
    saved(db);a=assign(db,'o1',11,10);assign(db,'o1',12,10)
    parcel(db,'o1','first',[a]);parcel(db,'o1','again',[a],reship=True,date='2026-09-12')
    decision={'reason':'ordinary','evidence':'核对补发工单'}
    result=snapshot(db,decisions={'again':decision});db.commit()
    app,ns=routes(db)
    ns['_calc_partner_recon_detail']=lambda *a,**k:{'total_orders':1,'total_gross_pln':225,'total_net_pln':200,'actual_cost_pln':80}
    ns['_compute_statement_split']=lambda *a,**k:(80,60,60)
    ns['_lookup_partner_rate']=lambda *a:(2,'2026-09')
    ns['_build_product_source_statement']=lambda *a,**k:result
    ns['_snapshot_statement_orders']=lambda *a:1
    response=app.test_client().post('/api/reconciliation/statements/generate',json={'partner_id':1,'year':2026,'month':9,
        'mode':'actual','reship_decisions':{'again':decision}})
    assert response.status_code==200,response.json
    record=db.execute("SELECT id,data FROM rec_objects WHERE kind='legacy_shipment_classification'").fetchone()
    assert json.loads(record['data'])['evidence']=='核对补发工单'
    stored=read_version(db,59,details='fees')
    assert next(row for row in stored['service_lines'] if row['shipment_id']=='again')['classification_id']==record['id']


def test_service_only_month_can_generate_without_new_sales_orders(db):
    saved(db);db.execute("UPDATE orders SET date_created='2026-08-20'")
    a=assign(db,'o1',11,10);assign(db,'o1',12,10);parcel(db,'o1','late-original',[a])
    value=snapshot(db);db.commit()
    assert value['source_lines']==[] and value['settlement']['packing_fee_cny']==20
    app,ns=routes(db)
    ns['_calc_partner_recon_detail']=lambda *a,**k:{'total_orders':0,'total_gross_pln':0,'total_net_pln':0,'actual_cost_pln':0}
    ns['_compute_statement_split']=lambda *a,**k:(0,0,0)
    ns['_lookup_partner_rate']=lambda *a:(2,'2026-09')
    ns['_build_product_source_statement']=lambda *a,**k:value
    ns['_snapshot_statement_orders']=lambda *a:0
    response=app.test_client().post('/api/reconciliation/statements/generate',json={'partner_id':1,'year':2026,'month':9,'mode':'actual'})
    assert response.status_code==200,response.json
    assert response.json['settlement']['packing_fee_cny']==20
    assert db.execute('SELECT total_orders FROM reconciliation_statements WHERE id=59').fetchone()[0]==0
